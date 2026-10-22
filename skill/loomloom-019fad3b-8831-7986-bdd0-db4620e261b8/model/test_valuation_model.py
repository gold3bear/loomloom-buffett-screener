#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Tests for the deterministic three-statement + DCF model.

Run with either interpreter:

    python3 -m unittest discover -s model -p "test_*.py"
    /path/to/bundled/python3 -m unittest discover -s model -p "test_*.py"
"""

from __future__ import annotations

import copy
import json
import os
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import valuation_model as vm  # noqa: E402


def load_fixture(name: str) -> dict:
    with open(os.path.join(HERE, "fixtures", name), "r", encoding="utf-8") as handle:
        return json.load(handle)


PACKETS = load_fixture("aapl-packets.json")
ASSUMPTIONS = load_fixture("aapl-assumptions.json")


def build(packets=None, assumptions=None) -> dict:
    return vm.build_model(json.dumps(packets or PACKETS), json.dumps(assumptions or ASSUMPTIONS))


class GoldenCaseTests(unittest.TestCase):
    def test_model_is_ready_and_has_three_scenarios(self):
        model = build()
        self.assertTrue(model["model_ready"], model["errors"])
        self.assertEqual(model["ticker"], "AAPL")
        self.assertEqual(set(model["scenarios"]), {"bear", "base", "bull"})
        self.assertEqual(model["baseScenario"], "base")
        self.assertEqual(len(model["scenarios"]["base"]["rows"]), 5)
        self.assertGreater(model["scenarios"]["base"]["summary"]["valuePerShare"], 0)

    def test_balance_sheet_balances_every_year(self):
        model = build()
        for name, result in model["scenarios"].items():
            for row in result["rows"]:
                self.assertAlmostEqual(row["totalAssets"], row["totalLiabilitiesAndEquity"], places=4, msg="%s FY+%d" % (name, row["year"]))
                self.assertAlmostEqual(row["balanceCheck"], 0.0, places=4, msg="%s FY+%d" % (name, row["year"]))
            self.assertTrue(result["integrity"]["passed"], result["checks"])

    def test_all_error_severity_checks_pass(self):
        model = build()
        for result in model["scenarios"].values():
            for check in result["checks"]:
                if check["severity"] == "error":
                    self.assertTrue(check["passed"], check)

    def test_scenario_ordering_is_monotonic(self):
        model = build()
        bear = model["scenarios"]["bear"]["summary"]["valuePerShare"]
        base = model["scenarios"]["base"]["summary"]["valuePerShare"]
        bull = model["scenarios"]["bull"]["summary"]["valuePerShare"]
        self.assertLess(bear, base)
        self.assertLess(base, bull)

    def test_sensitivity_grid_shape(self):
        model = build()
        sensitivity = model["sensitivity"]
        self.assertEqual(len(sensitivity["waccOffsets"]), 5)
        self.assertEqual(len(sensitivity["valuesPerShare"]), 5)
        for row in sensitivity["valuesPerShare"]:
            self.assertEqual(len(row), 5)
        center = sensitivity["valuesPerShare"][2][2]
        self.assertAlmostEqual(center, model["scenarios"]["base"]["summary"]["valuePerShare"], places=4)

    def test_reverse_dcf_reproduces_the_current_price(self):
        model = build()
        reverse = model["reverseDcf"]
        self.assertTrue(reverse["solved"], reverse)
        assumptions = copy.deepcopy(ASSUMPTIONS)
        assumptions["scenarios"]["base"]["revenueGrowth"] = reverse["impliedRevenueGrowth"]
        rebuilt = build(assumptions=assumptions)
        implied_value = rebuilt["scenarios"]["base"]["summary"]["valuePerShare"]
        self.assertAlmostEqual(implied_value, PACKETS["financialPacket"]["marketSnapshot"]["price"], delta=0.02)

    def test_model_invariants_pass(self):
        model = build()
        checks = model["modelChecks"]
        self.assertEqual(
            {check["name"] for check in checks},
            {"sensitivityCenterMatchesBase", "sensitivityWaccMonotonic", "sensitivityTerminalMonotonic", "reverseDcfReproducesPrice"},
        )
        for check in checks:
            self.assertTrue(check["passed"], check)
        self.assertTrue(model["reverseDcf"]["solved"])
        self.assertAlmostEqual(model["reverseDcf"]["impliedValuePerShare"], PACKETS["financialPacket"]["marketSnapshot"]["price"], delta=0.01)

    def test_deterministic_output(self):
        first = json.dumps(build(), sort_keys=True)
        second = json.dumps(build(), sort_keys=True)
        self.assertEqual(first, second)

    def test_exit_multiple_terminal_method(self):
        assumptions = copy.deepcopy(ASSUMPTIONS)
        assumptions["scenarios"]["base"]["terminalMethod"] = "exitMultiple"
        assumptions["scenarios"]["base"]["exitMultiple"] = 14.0
        assumptions["scenarios"]["base"].pop("terminalGrowth", None)
        model = build(assumptions=assumptions)
        self.assertTrue(model["model_ready"], model["errors"])
        summary = model["scenarios"]["base"]["summary"]
        self.assertEqual(summary["terminalMethod"], "exitMultiple")
        self.assertIn("EBIT x 14", summary["terminalAssumption"])

    def test_nwc_mismatch_is_surfaced(self):
        model = build()
        codes = {warning["code"] for warning in model["warnings"]}
        self.assertIn("NWC_DRIVER_MISMATCH", codes)


class FailureModeTests(unittest.TestCase):
    def test_missing_driver_blocks_the_model(self):
        assumptions = copy.deepcopy(ASSUMPTIONS)
        assumptions["scenarios"]["base"].pop("wacc")
        model = build(assumptions=assumptions)
        self.assertFalse(model["model_ready"])
        self.assertIn("scenarios.base.wacc", model["missingInputs"])
        self.assertTrue(any(error["code"] == vm.MISSING for error in model["errors"]))

    def test_unbalanced_opening_balance_sheet_blocks_the_model(self):
        packets = copy.deepcopy(PACKETS)
        packets["financialPacket"]["balanceSheet"]["otherAssets"] = 100000
        model = build(packets=packets)
        self.assertFalse(model["model_ready"])
        codes = {error["code"] for error in model["errors"]}
        self.assertIn("OPENING_BALANCE_SHEET_DOES_NOT_BALANCE", codes)

    def test_wacc_at_or_below_terminal_growth_blocks_the_model(self):
        assumptions = copy.deepcopy(ASSUMPTIONS)
        assumptions["scenarios"]["base"]["terminalGrowth"] = 0.09
        model = build(assumptions=assumptions)
        self.assertFalse(model["model_ready"])
        self.assertTrue(any(error["code"] == vm.INVALID for error in model["errors"]))

    def test_invalid_packet_json_is_reported(self):
        model = vm.build_model("{not json", json.dumps(ASSUMPTIONS))
        self.assertFalse(model["model_ready"])
        self.assertEqual(model["errors"][0]["code"], "PACKET_NOT_JSON")

    def test_empty_assumptions_is_reported(self):
        model = vm.build_model(json.dumps(PACKETS), "")
        self.assertFalse(model["model_ready"])
        self.assertEqual(model["errors"][0]["code"], "ASSUMPTIONS_NOT_JSON")

    def test_shares_are_required(self):
        packets = copy.deepcopy(PACKETS)
        for key in ("sharesOutstanding",):
            packets["financialPacket"]["balanceSheet"].pop(key, None)
            packets["financialPacket"]["marketSnapshot"].pop(key, None)
        assumptions = copy.deepcopy(ASSUMPTIONS)
        model = build(packets=packets, assumptions=assumptions)
        self.assertFalse(model["model_ready"])
        self.assertTrue(any("shares" in issue["path"] for issue in model["errors"]))


class WorkbookTests(unittest.TestCase):
    def test_main_writes_json_and_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            previous = os.getcwd()
            os.chdir(directory)
            try:
                result = vm.main({"tickerPacket": json.dumps(PACKETS), "assumptions": json.dumps(ASSUMPTIONS)})
                self.assertTrue(os.path.exists(result["workbook"]))
                self.assertEqual(result["model"]["schemaVersion"], vm.MODEL_SCHEMA)
            finally:
                os.chdir(previous)

    def test_workbook_opens_and_contains_expected_sheets(self):
        try:
            import openpyxl  # type: ignore
        except ImportError:
            self.skipTest("openpyxl is not installed in this interpreter")
        model = build()
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "model.xlsx")
            vm.write_xlsx(path, model)
            workbook = openpyxl.load_workbook(path)
            names = set(workbook.sheetnames)
            for expected in ("Cover", "Assumptions", "Historical", "P&L base", "BS base", "CF base", "DCF base", "Sensitivity", "Checks"):
                self.assertIn(expected, names)
            dcf = workbook["DCF base"]
            labels = [dcf.cell(row=row, column=1).value for row in range(1, dcf.max_row + 1)]
            self.assertIn("Margin of safety", labels)
            checks = workbook["Checks"]
            check_names = [checks.cell(row=row, column=2).value for row in range(1, checks.max_row + 1)]
            self.assertIn("sensitivityCenterMatchesBase", check_names)
            self.assertIn("reverseDcfReproducesPrice", check_names)

    def test_workbook_artifact_is_byte_stable(self):
        model = build()
        with tempfile.TemporaryDirectory() as directory:
            first = os.path.join(directory, "a.xlsx")
            second = os.path.join(directory, "b.xlsx")
            vm.write_xlsx(first, model)
            vm.write_xlsx(second, model)
            with open(first, "rb") as handle:
                first_bytes = handle.read()
            with open(second, "rb") as handle:
                second_bytes = handle.read()
            self.assertEqual(first_bytes, second_bytes)


if __name__ == "__main__":
    unittest.main(verbosity=2)
